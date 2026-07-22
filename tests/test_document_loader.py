"""Pruebas de la capa de ingesta de documentos."""

from __future__ import annotations

import json
from pathlib import Path

import fitz
import pytest
from docx import Document as DocxDocument
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from pptx import Presentation

from src.ingestion.document_loader import (
    DocumentLoadError,
    SUPPORTED_EXTENSIONS,
    UnsupportedDocumentTypeError,
    load_document,
    load_documents,
)


def _write_pdf(path: Path, text: str) -> None:
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), text)
    pdf.save(path)
    pdf.close()


def _write_docx(path: Path) -> None:
    document = DocxDocument()
    document.add_paragraph("Párrafo corporativo")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Área"
    table.cell(0, 1).text = "Ventas"
    document.save(path)


def _write_xlsx(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Resumen"
    worksheet.append(["Producto", "Cantidad"])
    worksheet.append(["Café", 3])
    workbook.save(path)


def _write_pptx(path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.add_textbox(0, 0, 1000000, 1000000).text_frame.text = "Informe mensual"
    presentation.save(path)


def test_load_markdown_and_common_metadata(tmp_path: Path) -> None:
    path = tmp_path / "nota.MD"
    path.write_text("# Política\nContenido", encoding="utf-8")

    documents = load_document(path)

    assert documents[0].page_content == "# Política\nContenido"
    assert documents[0].metadata == {"source": "nota.MD", "file_name": "nota.MD", "file_type": "md"}


def test_load_markdown_with_utf8_bom(tmp_path: Path) -> None:
    path = tmp_path / "bom.md"
    path.write_text("Contenido con BOM", encoding="utf-8-sig")

    assert load_document(path)[0].page_content == "Contenido con BOM"


def test_load_csv_preserves_rows(tmp_path: Path) -> None:
    path = tmp_path / "datos.csv"
    path.write_text('nombre,detalle\nAna,"uno, dos"\n', encoding="utf-8")

    document = load_document(path)[0]

    assert document.page_content == 'nombre,detalle\nAna,"uno, dos"'


def test_load_json_normalizes_unicode(tmp_path: Path) -> None:
    path = tmp_path / "datos.json"
    path.write_text(json.dumps({"nombre": "José", "activo": True}), encoding="utf-8")

    document = load_document(path)[0]

    assert document.page_content == '{\n  "activo": true,\n  "nombre": "José"\n}'


def test_load_html_excludes_script_and_style(tmp_path: Path) -> None:
    path = tmp_path / "pagina.html"
    path.write_text("<style>oculto</style><h1>Título</h1><script>secreto()</script><p>Visible</p>", encoding="utf-8")

    document = load_document(path)[0]

    assert document.page_content == "Título\nVisible"
    assert "oculto" not in document.page_content
    assert "secreto" not in document.page_content


def test_load_htm_with_uppercase_extension(tmp_path: Path) -> None:
    path = tmp_path / "pagina.HTM"
    path.write_text("<p>Contenido HTM</p>", encoding="utf-8")

    document = load_document(path)[0]

    assert document.page_content == "Contenido HTM"
    assert document.metadata["file_type"] == "htm"


def test_load_pdf_has_page_metadata(tmp_path: Path) -> None:
    path = tmp_path / "informe.pdf"
    _write_pdf(path, "Reporte PDF")

    document = load_document(path)[0]

    assert document.page_content == "Reporte PDF"
    assert document.metadata["page"] == 1


def test_load_pdf_returns_one_document_per_nonempty_page(tmp_path: Path) -> None:
    path = tmp_path / "multipagina.pdf"
    pdf = fitz.open()
    for text in ("Primera página", "", "Tercera página"):
        page = pdf.new_page()
        if text:
            page.insert_text((72, 72), text)
    pdf.save(path)
    pdf.close()

    documents = load_document(path)

    assert [document.page_content for document in documents] == ["Primera página", "Tercera página"]
    assert [document.metadata["page"] for document in documents] == [1, 3]


def test_load_docx_includes_paragraphs_and_tables(tmp_path: Path) -> None:
    path = tmp_path / "informe.docx"
    _write_docx(path)

    document = load_document(path)[0]

    assert "Párrafo corporativo" in document.page_content
    assert "Área | Ventas" in document.page_content


def test_load_xlsx_has_sheet_metadata(tmp_path: Path) -> None:
    path = tmp_path / "datos.xlsx"
    _write_xlsx(path)

    document = load_document(path)[0]

    assert document.metadata["sheet"] == "Resumen"
    assert "Producto | Cantidad" in document.page_content


def test_load_xlsx_ignores_styled_cells_without_values(tmp_path: Path) -> None:
    path = tmp_path / "vacio.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    fill = PatternFill(fill_type="solid", fgColor="FFFFFF")
    worksheet["A1"].fill = fill
    worksheet["B1"].fill = fill
    workbook.save(path)

    assert load_document(path) == []


def test_load_pptx_has_slide_metadata(tmp_path: Path) -> None:
    path = tmp_path / "presentacion.pptx"
    _write_pptx(path)

    document = load_document(path)[0]

    assert document.page_content == "Informe mensual"
    assert document.metadata["slide"] == 1


def test_load_pptx_returns_one_document_per_nonempty_slide(tmp_path: Path) -> None:
    path = tmp_path / "multipresentacion.pptx"
    presentation = Presentation()
    for text in ("Primera diapositiva", "", "Tercera diapositiva"):
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        if text:
            slide.shapes.add_textbox(0, 0, 1000000, 1000000).text_frame.text = text
    presentation.save(path)

    documents = load_document(path)

    assert [document.page_content for document in documents] == ["Primera diapositiva", "Tercera diapositiva"]
    assert [document.metadata["slide"] for document in documents] == [1, 3]


def test_rejects_unsupported_extension(tmp_path: Path) -> None:
    path = tmp_path / "datos.txt"
    path.write_text("texto", encoding="utf-8")

    with pytest.raises(UnsupportedDocumentTypeError, match="no admitida"):
        load_document(path)


def test_rejects_missing_file(tmp_path: Path) -> None:
    with pytest.raises(FileNotFoundError, match="no existe"):
        load_document(tmp_path / "ausente.md")


def test_rejects_directory_passed_as_document(tmp_path: Path) -> None:
    with pytest.raises(IsADirectoryError, match="no es un archivo"):
        load_document(tmp_path)


def test_rejects_missing_directory(tmp_path: Path) -> None:
    with pytest.raises(FileNotFoundError, match="directorio no existe"):
        load_documents(tmp_path / "ausente")


def test_rejects_file_passed_as_directory(tmp_path: Path) -> None:
    path = tmp_path / "archivo.md"
    path.write_text("contenido", encoding="utf-8")

    with pytest.raises(NotADirectoryError, match="no es un directorio"):
        load_documents(path)


def test_wraps_corrupt_supported_file_and_preserves_cause(tmp_path: Path) -> None:
    path = tmp_path / "corrupto.json"
    path.write_text("{", encoding="utf-8")

    with pytest.raises(DocumentLoadError, match="No se pudo leer") as error:
        load_document(path)

    assert error.value.__cause__ is not None
    assert isinstance(error.value.__cause__, json.JSONDecodeError)


def test_load_documents_recursively_and_in_order(tmp_path: Path) -> None:
    nested = tmp_path / "sub"
    nested.mkdir()
    (nested / "b.md").write_text("segundo", encoding="utf-8")
    (tmp_path / "a.md").write_text("primero", encoding="utf-8")
    (tmp_path / "ignorado.txt").write_text("ignorado", encoding="utf-8")

    documents = load_documents(tmp_path)

    assert [document.page_content for document in documents] == ["primero", "segundo"]
    assert [document.metadata["file_name"] for document in documents] == ["a.md", "b.md"]


def test_empty_content_is_excluded(tmp_path: Path) -> None:
    path = tmp_path / "vacio.md"
    path.write_text(" \n\t", encoding="utf-8")

    assert load_document(path) == []


def test_source_is_never_absolute_and_uses_portable_separators(tmp_path: Path) -> None:
    path = tmp_path / "subcarpeta" / "archivo.md"
    path.parent.mkdir()
    path.write_text("contenido", encoding="utf-8")

    source = load_document(path)[0].metadata["source"]

    assert source == "archivo.md"
    assert not Path(source).is_absolute()
    assert "\\" not in source


def test_supported_extensions_are_public() -> None:
    assert SUPPORTED_EXTENSIONS == {".pdf", ".docx", ".xlsx", ".pptx", ".md", ".csv", ".json", ".html", ".htm"}
