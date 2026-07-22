"""Carga documentos corporativos en objetos de LangChain."""

from __future__ import annotations

import csv
import json
from collections.abc import Callable, Iterable
from io import StringIO
from pathlib import Path

import fitz
from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from langchain_core.documents import Document
from openpyxl import load_workbook
from pptx import Presentation

SUPPORTED_EXTENSIONS = frozenset({".pdf", ".docx", ".xlsx", ".pptx", ".md", ".csv", ".json", ".html", ".htm"})


class DocumentLoadError(Exception):
    """Indica que no fue posible leer un documento admitido."""


class UnsupportedDocumentTypeError(DocumentLoadError):
    """Indica que un archivo no tiene una extensión admitida."""


def _portable_source(path: Path) -> str:
    """Devuelve una fuente sin rutas absolutas."""
    return path.as_posix() if not path.is_absolute() else path.name


def _metadata(path: Path, **extra: str | int) -> dict[str, str | int]:
    """Construye metadatos comunes sin exponer rutas absolutas."""
    return {
        "source": _portable_source(path),
        "file_name": path.name,
        "file_type": path.suffix.lower().lstrip("."),
        **extra,
    }


def _document(path: Path, content: str, **extra: str | int) -> Document | None:
    """Crea un documento solo si el contenido contiene texto."""
    normalized = content.strip()
    if not normalized:
        return None
    return Document(page_content=normalized, metadata=_metadata(path, **extra))


def _documents(path: Path, units: Iterable[tuple[str, dict[str, str | int]]]) -> list[Document]:
    """Convierte unidades de texto en documentos no vacíos."""
    return [document for content, extra in units if (document := _document(path, content, **extra)) is not None]


def _load_pdf(path: Path) -> list[Document]:
    with fitz.open(path) as pdf:
        return _documents(path, ((page.get_text("text"), {"page": number}) for number, page in enumerate(pdf, start=1)))


def _load_docx(path: Path) -> list[Document]:
    document = DocxDocument(path)
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    rows = [" | ".join(cell.text.strip() for cell in row.cells).strip() for table in document.tables for row in table.rows]
    content = "\n".join([*paragraphs, *(row for row in rows if row)])
    return _documents(path, [(content, {})])


def _cell_value(value: object) -> str:
    """Convierte valores de hoja a texto estable."""
    return "" if value is None else str(value)


def _has_cell_value(value: object) -> bool:
    """Indica si una celda aporta contenido a su fila."""
    return value is not None and (not isinstance(value, str) or bool(value.strip()))


def _load_xlsx(path: Path) -> list[Document]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        units: list[tuple[str, dict[str, str | int]]] = []
        for worksheet in workbook.worksheets:
            rows = [
                " | ".join(_cell_value(value) for value in row).rstrip()
                for row in worksheet.iter_rows(values_only=True)
                if any(_has_cell_value(value) for value in row)
            ]
            content = "\n".join(rows)
            units.append((content, {"sheet": worksheet.title}))
        return _documents(path, units)
    finally:
        workbook.close()


def _load_pptx(path: Path) -> list[Document]:
    presentation = Presentation(path)
    units = []
    for number, slide in enumerate(presentation.slides, start=1):
        text = "\n".join(shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False))
        units.append((text, {"slide": number}))
    return _documents(path, units)


def _load_text(path: Path) -> list[Document]:
    return _documents(path, [(path.read_text(encoding="utf-8-sig"), {})])


def _load_csv(path: Path) -> list[Document]:
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        rows = list(csv.reader(file))
    output = StringIO()
    writer = csv.writer(output, lineterminator="\n")
    writer.writerows(rows)
    return _documents(path, [(output.getvalue(), {})])


def _load_json(path: Path) -> list[Document]:
    with path.open("r", encoding="utf-8-sig") as file:
        content = json.load(file)
    return _documents(path, [(json.dumps(content, ensure_ascii=False, indent=2, sort_keys=True), {})])


def _load_html(path: Path) -> list[Document]:
    soup = BeautifulSoup(path.read_text(encoding="utf-8-sig"), "html.parser")
    for element in soup(["script", "style"]):
        element.decompose()
    return _documents(path, [(soup.get_text("\n", strip=True), {})])


_LOADERS: dict[str, Callable[[Path], list[Document]]] = {
    ".pdf": _load_pdf,
    ".docx": _load_docx,
    ".xlsx": _load_xlsx,
    ".pptx": _load_pptx,
    ".md": _load_text,
    ".csv": _load_csv,
    ".json": _load_json,
    ".html": _load_html,
    ".htm": _load_html,
}


def load_document(path: str | Path) -> list[Document]:
    """Carga un archivo admitido y devuelve sus unidades no vacías."""
    file_path = Path(path)
    if not file_path.exists():
        raise FileNotFoundError(f"El archivo no existe: {file_path}")
    if not file_path.is_file():
        raise IsADirectoryError(f"La ruta no es un archivo: {file_path}")

    extension = file_path.suffix.lower()
    loader = _LOADERS.get(extension)
    if loader is None:
        raise UnsupportedDocumentTypeError(f"Extensión no admitida: {file_path.suffix or '(sin extensión)'}")

    try:
        return loader(file_path)
    except DocumentLoadError:
        raise
    except Exception as error:
        raise DocumentLoadError(f"No se pudo leer el documento '{file_path.name}': {error}") from error


def load_documents(directory: str | Path) -> list[Document]:
    """Carga recursivamente los archivos admitidos de un directorio."""
    directory_path = Path(directory)
    if not directory_path.exists():
        raise FileNotFoundError(f"El directorio no existe: {directory_path}")
    if not directory_path.is_dir():
        raise NotADirectoryError(f"La ruta no es un directorio: {directory_path}")

    paths = sorted(
        (path for path in directory_path.rglob("*") if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS),
        key=lambda path: path.as_posix().lower(),
    )
    return [document for path in paths for document in load_document(path)]
